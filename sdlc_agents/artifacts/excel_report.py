"""Excel validation report writer (openpyxl).

Takes the structured validation result and produces a multi-sheet ``.xlsx``
workbook entirely in memory (returned as bytes), so nothing touches the network
and the caller controls where/how it is stored.
"""

from __future__ import annotations

import io
from typing import Dict, List


# Colours for PASS / PARTIAL / FAIL verdicts.
_FILLS = {
    "PASS": "C6EFCE",
    "PARTIAL": "FFEB9C",
    "FAIL": "FFC7CE",
    "N/A": "D9D9D9",
}
_FONTS = {
    "PASS": "006100",
    "PARTIAL": "9C6500",
    "FAIL": "9C0006",
    "N/A": "3F3F3F",
}


def build_validation_workbook(report: Dict) -> bytes:
    """Render the validation ``report`` dict into xlsx bytes."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    title_font = Font(bold=True, size=14)
    wrap = Alignment(wrap_text=True, vertical="top")

    def verdict_cell(ws, row, col, verdict):
        cell = ws.cell(row=row, column=col, value=verdict)
        v = verdict.upper()
        if v in _FILLS:
            cell.fill = PatternFill("solid", fgColor=_FILLS[v])
            cell.font = Font(bold=True, color=_FONTS[v])
        return cell

    def autosize(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def add_table(title, headers, rows, verdict_col, widths):
        ws = wb.create_sheet(title=title)
        # Title row (row 1)
        tcell = ws.cell(row=1, column=1, value=title)
        tcell.font = title_font
        # Header row (row 2)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
        # Data rows
        for r, row in enumerate(rows, start=3):
            for c, val in enumerate(row, start=1):
                if c == verdict_col:
                    verdict_cell(ws, r, c, str(val))
                else:
                    cell = ws.cell(row=r, column=c, value=val)
                    cell.alignment = wrap
        autosize(ws, widths)
        ws.freeze_panes = "A3"
        return ws

    # ---- Summary sheet ---------------------------------------------------
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Validation Report"
    summary["A1"].font = Font(bold=True, size=16)
    meta = report.get("meta", {})
    rows = [
        ("Project", meta.get("title", "")),
        ("Run ID", meta.get("run_id", "")),
        ("Generated", meta.get("generated", "")),
        ("Backend", meta.get("backend", "")),
        ("Overall Verdict", report.get("overall", "")),
        ("Functional Requirements", meta.get("fr_count", 0)),
        ("Non-Functional Requirements", meta.get("nfr_count", 0)),
        ("Tests Passed", meta.get("tests_passed", 0)),
        ("Tests Failed", meta.get("tests_failed", 0)),
        ("Tests Total", meta.get("tests_total", 0)),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        summary.cell(row=i, column=1, value=k).font = Font(bold=True)
        if k == "Overall Verdict":
            verdict_cell(summary, i, 2, str(v))
        else:
            summary.cell(row=i, column=2, value=v)
    autosize(summary, [32, 60])

    # ---- Detail sheets ---------------------------------------------------
    add_table(
        "Functional",
        ["ID", "Requirement", "Verification Method", "Verdict", "Notes"],
        [[r["id"], r["requirement"], r["method"], r["verdict"], r["notes"]] for r in report.get("functional", [])],
        verdict_col=4,
        widths=[10, 55, 26, 12, 45],
    )
    add_table(
        "NonFunctional",
        ["ID", "Requirement", "Tactic / Check", "Verdict", "Notes"],
        [[r["id"], r["requirement"], r["method"], r["verdict"], r["notes"]] for r in report.get("nonfunctional", [])],
        verdict_col=4,
        widths=[10, 55, 30, 12, 45],
    )
    add_table(
        "Architecture",
        ["ID", "Check", "Verdict", "Notes"],
        [[r["id"], r["check"], r["verdict"], r["notes"]] for r in report.get("architecture", [])],
        verdict_col=3,
        widths=[10, 45, 12, 55],
    )
    add_table(
        "TestResults",
        ["Suite", "Total", "Passed", "Failed", "Verdict"],
        [[r["suite"], r["total"], r["passed"], r["failed"], r["verdict"]] for r in report.get("tests", [])],
        verdict_col=5,
        widths=[26, 10, 10, 10, 12],
    )
    add_table(
        "CodeQuality",
        ["Metric", "Value", "Verdict", "Notes"],
        [[r["metric"], r["value"], r["verdict"], r["notes"]] for r in report.get("quality", [])],
        verdict_col=3,
        widths=[30, 18, 12, 50],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
