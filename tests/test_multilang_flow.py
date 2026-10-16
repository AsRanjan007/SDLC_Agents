"""End-to-end orchestrator tests across languages and agent configuration.

Test execution is disabled (Developer/Validation ``run_tests = False``) so these
tests are deterministic regardless of which language toolchains are installed on
the host. Native execution is covered separately where a toolchain is present.
"""

import io

import pytest
from openpyxl import load_workbook

from sdlc_agents.agents import Orchestrator
from sdlc_agents.config import AppConfig
from sdlc_agents.languages import LANGUAGE_ORDER, get_language


def _cfg(language, framework="auto"):
    cfg = AppConfig()
    cfg.backend = "offline"
    cfg.audio.enabled = False
    cfg.target_language = language
    cfg.test_framework = framework
    cfg.agents.developer.run_tests = False
    cfg.agents.validation.run_tests = False
    return cfg


@pytest.mark.parametrize("language", list(LANGUAGE_ORDER))
def test_flow_produces_language_artifacts(language):
    result = Orchestrator(_cfg(language)).run(
        "Read a number from input; print Even if divisible by 2, else Odd.",
        "an integer",
        "Even/Odd output plus a validation report",
    )
    assert len(result.results) == 3
    paths = {a.path for r in result.results for a in r.artifacts}

    assert "docs/architecture/Software_Architecture_Document.md" in paths
    assert "docs/design/Component_Design.md" in paths
    assert "docs/reports/Developer_Build_Report.md" in paths
    assert "docs/reports/Validation_Report.xlsx" in paths

    ext = get_language(language).ext
    assert any(p.endswith(ext) for p in paths)
    manifest = get_language(language).manifest_name
    assert any(p.split("/")[-1] == manifest for p in paths)


@pytest.mark.parametrize("language", list(LANGUAGE_ORDER))
def test_developer_metrics_report_language_and_framework(language):
    result = Orchestrator(_cfg(language)).run("Process and report records.", "a record", "a report")
    dev = next(r for r in result.results if r.agent_id == "developer")
    assert dev.metrics["language"] == get_language(language).display
    assert dev.metrics["test_framework"] == get_language(language).default_framework.display
    # Execution disabled -> deterministic 'skipped' status.
    assert dev.metrics["test_status"] == "skipped"


def test_validation_excel_readable_for_rust():
    result = Orchestrator(_cfg("rust")).run("Store and fetch documents.", "a document", "stored id")
    val = next(r for r in result.results if r.agent_id == "validation_engineer")
    xlsx = next(a for a in val.artifacts if a.path.endswith(".xlsx"))
    assert xlsx.data and xlsx.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(xlsx.data))
    assert "Functional" in wb.sheetnames
    assert (result.workspace / "docs/reports/Validation_Report.xlsx").exists()


def test_architect_config_toggles_sections():
    cfg = _cfg("python")
    cfg.agents.architect.include_risks = False
    cfg.agents.architect.include_adrs = False
    result = Orchestrator(cfg).run("Build a small service.", "input", "output")
    arch = next(r for r in result.results if r.agent_id == "architect")
    doc = arch.artifacts[0].content
    assert "Risks & Mitigations" not in doc
    assert "Architecture Decisions" not in doc
    # Core sections remain.
    assert "Software Architecture Document" in doc
    assert "Traceability" in doc


def test_developer_config_skips_integration_tests():
    cfg = _cfg("rust")
    cfg.agents.developer.generate_integration_tests = False
    result = Orchestrator(cfg).run("Build a small service.", "input", "output")
    paths = {a.path for r in result.results for a in r.artifacts}
    assert "tests/integration.rs" not in paths
    # Source still present.
    assert "src/lib.rs" in paths


def test_validation_config_can_disable_excel():
    cfg = _cfg("python")
    cfg.agents.validation.excel_report = False
    result = Orchestrator(cfg).run("Build a small service.", "input", "output")
    val = next(r for r in result.results if r.agent_id == "validation_engineer")
    assert not any(a.path.endswith(".xlsx") for a in val.artifacts)
    assert any(a.path.endswith("Validation_Report.md") for a in val.artifacts)
