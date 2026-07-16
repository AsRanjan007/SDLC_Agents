from sdlc_agents.artifacts.excel_report import build_validation_workbook


def _sample_report():
    return {
        "meta": {"title": "Demo", "run_id": "r1", "generated": "now", "backend": "offline",
                 "fr_count": 1, "nfr_count": 1, "tests_passed": 2, "tests_failed": 0, "tests_total": 2},
        "overall": "PASS",
        "functional": [{"id": "FR-1", "requirement": "do x", "method": "test", "verdict": "PASS", "notes": "ok"}],
        "nonfunctional": [{"id": "NFR-1", "requirement": "offline", "method": "scan", "verdict": "PASS", "notes": "ok"}],
        "architecture": [{"id": "AR-1", "check": "SAD present", "verdict": "PASS", "notes": "ok"}],
        "tests": [{"suite": "unit", "total": 2, "passed": 2, "failed": 0, "verdict": "PASS"}],
        "quality": [{"metric": "loc", "value": 100, "verdict": "N/A", "notes": "ok"}],
    }


def test_workbook_is_valid_xlsx_bytes():
    data = build_validation_workbook(_sample_report())
    assert isinstance(data, bytes)
    # xlsx is a zip archive -> starts with PK.
    assert data[:2] == b"PK"


def test_workbook_has_expected_sheets():
    import io

    from openpyxl import load_workbook

    data = build_validation_workbook(_sample_report())
    wb = load_workbook(io.BytesIO(data))
    for sheet in ("Summary", "Functional", "NonFunctional", "Architecture", "TestResults", "CodeQuality"):
        assert sheet in wb.sheetnames
