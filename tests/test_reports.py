"""Tests for timestamps, endpoint labels and the DevMetrics Excel sheet."""

import io
import os
import tempfile

from openpyxl import load_workbook

from sdlc_agents import util
from sdlc_agents.agents import Orchestrator
from sdlc_agents.config import AppConfig, OpenAICompatConfig, ProxyConfig
from sdlc_agents.llm.openai_compat_backend import OpenAICompatBackend


# ---- timestamps -----------------------------------------------------------


def test_now_str_is_timezone_labelled():
    s = util.now_str("UTC")
    assert "UTC" in s
    # Shape: YYYY-MM-DD HH:MM:SS <zone>
    assert len(s) >= 20 and s[4] == "-" and s[13] == ":"


def test_now_str_honours_explicit_zone():
    kolkata = util.now_str("Asia/Kolkata")
    # IST is UTC+5:30 — label is either the name or the offset.
    assert ("IST" in kolkata) or ("+0530" in kolkata) or ("+05" in kolkata)


def test_apply_timezone_sets_default(monkeypatch):
    util.apply_timezone("UTC")
    try:
        assert "UTC" in util.now_str()
    finally:
        util.apply_timezone("")


def test_config_timezone_roundtrips():
    cfg = AppConfig()
    cfg.timezone = "America/New_York"
    assert AppConfig.from_dict(cfg.to_dict()).timezone == "America/New_York"


# ---- endpoint label -------------------------------------------------------


def test_openai_compat_label_uses_model_and_endpoint():
    cfg = OpenAICompatConfig(base_url="http://127.0.0.1:11434/v1", model="gemma4:e4b")
    backend = OpenAICompatBackend(cfg, ProxyConfig())
    assert backend.label == "gemma4:e4b @ 127.0.0.1:11434"
    assert backend.label != backend.name  # not the raw 'openai_compat' id


# ---- DevMetrics sheet + report timestamp ----------------------------------


def _run(tmp_tz=""):
    os.environ["SDLC_AGENTS_HOME"] = tempfile.mkdtemp(prefix="sdlc_rep_")
    cfg = AppConfig()
    cfg.backend = "offline"
    cfg.audio.enabled = False
    cfg.target_language = "python"
    cfg.timezone = tmp_tz
    return Orchestrator(cfg).run("Calculator for +,-,*,/.", "operands", "result")


def test_validation_workbook_has_devmetrics_sheet():
    result = _run()
    val = next(r for r in result.results if r.agent_id == "validation_engineer")
    xlsx = next(a for a in val.artifacts if a.path.endswith(".xlsx"))
    wb = load_workbook(io.BytesIO(xlsx.data))
    assert "DevMetrics" in wb.sheetnames
    ws = wb["DevMetrics"]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Code coverage" in text
    assert "Test status" in text


def test_reports_carry_timezoned_timestamp():
    result = _run(tmp_tz="UTC")
    dev = next(r for r in result.results if r.agent_id == "developer")
    report = next(a for a in dev.artifacts if a.path.endswith("Developer_Build_Report.md")).content
    assert "UTC" in report  # the build report timestamp is timezone-labelled

    arch = next(r for r in result.results if r.agent_id == "architect")
    sad = arch.artifacts[0].content
    assert "UTC" in sad  # offline SAD byline is timezone-labelled
