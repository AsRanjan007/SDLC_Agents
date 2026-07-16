import io

from openpyxl import load_workbook

from sdlc_agents.agents import Orchestrator
from sdlc_agents.config import AppConfig


def _offline_config():
    cfg = AppConfig()
    cfg.backend = "offline"
    cfg.audio.enabled = False  # no TTS in CI
    return cfg


def test_full_offline_flow_produces_all_artifacts():
    cfg = _offline_config()
    orch = Orchestrator(cfg)
    result = orch.run(
        instruction="Create and validate customer orders, then generate a report.",
        input_expectations="JSON order with id and items",
        output_expectations="Processed order plus a summary report",
    )

    assert len(result.results) == 3
    paths = {art.path for res in result.results for art in res.artifacts}
    assert "docs/architecture/Software_Architecture_Document.md" in paths
    assert "docs/design/Component_Design.md" in paths
    assert "docs/reports/Developer_Build_Report.md" in paths
    assert "docs/reports/Validation_Report.xlsx" in paths
    assert any(p.startswith("src/") and p.endswith(".py") for p in paths)
    assert any("tests/unit" in p for p in paths)
    assert any("tests/integration" in p for p in paths)


def test_generated_tests_actually_pass():
    cfg = _offline_config()
    result = Orchestrator(cfg).run(
        "Build a service that processes and reports on records.",
        "a record",
        "a report",
    )
    dev = next(r for r in result.results if r.agent_id == "developer")
    assert dev.metrics["test_status"] == "passed"
    assert dev.metrics["tests_passed"] >= 1
    assert dev.metrics["tests_failed"] == 0


def test_validation_excel_written_and_readable():
    cfg = _offline_config()
    result = Orchestrator(cfg).run("Store and fetch documents.", "a document", "stored id")
    val = next(r for r in result.results if r.agent_id == "validation_engineer")
    xlsx = next(a for a in val.artifacts if a.path.endswith(".xlsx"))
    assert xlsx.data and xlsx.data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(xlsx.data))
    assert "Functional" in wb.sheetnames
    # File also persisted to the local workspace.
    assert (result.workspace / "docs/reports/Validation_Report.xlsx").exists()


def test_offline_flow_makes_no_network_backend():
    cfg = _offline_config()
    orch = Orchestrator(cfg)
    assert orch.backend is None  # offline == no backend object, zero network
